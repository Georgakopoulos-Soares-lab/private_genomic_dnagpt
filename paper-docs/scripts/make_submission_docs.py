#!/usr/bin/env python3
"""Generate the Conflict of Interest and Graphical Abstract Text drafts as .docx.

The IEEE Author Portal accepts PDF, RTF, or MS Word for these two fields -- not plain text --
so the wording lives here rather than in a .txt file. Edit the strings below and rerun.
"""
from __future__ import annotations

import pathlib

import docx
from docx.shared import Pt, RGBColor

OUT = pathlib.Path(__file__).resolve().parents[1] / "submission"
DRAFT_RED = RGBColor(0xC0, 0x00, 0x00)


def base_doc() -> docx.Document:
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    return d


def draft_flag(d: docx.Document, text: str) -> None:
    run = d.add_paragraph().add_run(text)
    run.bold = True
    run.font.color.rgb = DRAFT_RED


def heading(d: docx.Document, text: str) -> None:
    run = d.add_paragraph().add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def conflict_of_interest() -> None:
    d = base_doc()
    draft_flag(d, "DRAFT — pending confirmation by all authors before submission.")
    heading(d, "Conflict of Interest Statement")
    d.add_paragraph('Manuscript: "Feasibility of Homomorphic Inference for a Genomic Foundation Model"')
    d.add_paragraph("Authors: Christos Galanopoulos, Kimon Antonios Provatas, Ilias Georgakopoulos-Soares")
    d.add_paragraph(
        "The authors declare that they have no known competing financial interests or personal "
        "relationships that could have appeared to influence the work reported in this paper."
    )
    d.add_paragraph().add_run(
        "Each author must confirm this statement applies to them, or amend it to disclose any "
        "conflict, before it is uploaded to the submission portal."
    ).italic = True
    d.save(OUT / "DNAGPT-FHE_conflict_of_interest_DRAFT.docx")


def graphical_abstract_text() -> None:
    d = base_doc()
    draft_flag(d, "DRAFT — author review recommended before submission.")
    heading(d, "Graphical Abstract Text")
    d.add_paragraph(
        "A client-assisted homomorphic-encryption protocol runs all twelve released transformer "
        "blocks of a genomic foundation model plus its classification head on an encrypted "
        "103-token sequence, matching the plaintext prediction while revealing no query genomic "
        "data to the compute provider — establishing feasibility, with practicality remaining "
        "an open, cost-bound question."
    )
    d.save(OUT / "DNAGPT-FHE_graphical_abstract_text.docx")


if __name__ == "__main__":
    conflict_of_interest()
    graphical_abstract_text()
    print(f"wrote {OUT / 'DNAGPT-FHE_conflict_of_interest_DRAFT.docx'}")
    print(f"wrote {OUT / 'DNAGPT-FHE_graphical_abstract_text.docx'}")
